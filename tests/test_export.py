from __future__ import annotations

import io

import pandas as pd
from openpyxl import load_workbook

from stock_metrics.export import metrics_workbook
from stock_metrics.metrics import compute_metrics


def test_metrics_workbook() -> None:
    base = pd.DataFrame({"revenue": [100e6, 120e6], "net_income": [10e6, 15e6]}, index=[2022, 2023])
    provenance = pd.DataFrame(
        {"concept": ["revenue"], "fiscal_year": [2023], "value": [120e6], "overridden": [False]}
    )
    data = compute_metrics(base)
    wb = load_workbook(io.BytesIO(metrics_workbook("TST", "Test Co", data, provenance)))
    assert wb.sheetnames == ["Metrics", "Sources"]
    ws = wb["Metrics"]
    assert ws["A1"].value.startswith("Test Co (TST)")
    assert [c.value for c in ws[3]][:3] == ["Metric", "FY2022", "FY2023"]
    assert ws["A4"].value == "Revenue"
    assert ws["C4"].value == 120e6
    assert ws["C4"].number_format == '#,##0.0,,"M"'
    net_margin_row = next(r for r in ws.iter_rows(min_row=4) if r[0].value == "Net Margin")
    assert net_margin_row[2].number_format == "0.0%"
    assert "overridden" not in [c.value for c in wb["Sources"][1]]
