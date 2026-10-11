"""Excel export of a company's metrics and source figures."""

from __future__ import annotations

import io

import pandas as pd
from openpyxl.utils import get_column_letter

from stock_metrics.catalog import METRICS, Unit

NUMBER_FORMATS = {
    Unit.CURRENCY: '#,##0.0,,"M"',
    Unit.PER_SHARE: "$#,##0.00",
    Unit.COUNT: '#,##0.0,,"M"',
    Unit.PERCENT: "0.0%",
    Unit.MULTIPLE: '0.0"x"',
    Unit.RATIO: "0.00",
}


def metrics_workbook(
    ticker: str, name: str, metrics: pd.DataFrame, provenance: pd.DataFrame
) -> bytes:
    """Workbook with a 'Metrics' sheet (metrics x fiscal years) and a 'Sources' sheet."""
    defs = [m for m in METRICS if m.key in metrics.columns]
    table = metrics[[m.key for m in defs]].T
    table.index = pd.Index([m.label for m in defs], name="Metric")
    table.columns = [f"FY{int(y)}" for y in table.columns]

    sources = provenance.drop(columns=["overridden"], errors="ignore")

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name="Metrics", startrow=2)
        sources.to_excel(writer, sheet_name="Sources", index=False)

        ws = writer.sheets["Metrics"]
        ws["A1"] = f"{name} ({ticker}) — currency and share counts in millions"
        ws.column_dimensions["A"].width = 34
        for col in range(2, table.shape[1] + 2):
            ws.column_dimensions[get_column_letter(col)].width = 12
        for row_offset, metric in enumerate(defs):
            row = 4 + row_offset  # title row, blank row, header row
            for col in range(2, table.shape[1] + 2):
                ws.cell(row=row, column=col).number_format = NUMBER_FORMATS[metric.unit]
        ws.freeze_panes = "B4"

        src = writer.sheets["Sources"]
        for col in range(1, sources.shape[1] + 1):
            src.column_dimensions[get_column_letter(col)].width = 18
    return buf.getvalue()
